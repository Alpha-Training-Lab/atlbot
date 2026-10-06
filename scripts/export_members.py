"""Export every member to a spreadsheet, with their ID photos. Owner only.

Run on the server, from the project root:
  .venv/bin/python scripts/export_members.py                  # everyone, with photos
  .venv/bin/python scripts/export_members.py --status active  # active members only
  .venv/bin/python scripts/export_members.py --no-photos      # spreadsheet only

Writes data/exports/atl_members_<time>.zip containing:
  members.xlsx   one row per member, every detail, a link to each photo
  ids/           the ID photos, downloaded from Telegram

Read-only on the database. The bot can keep running.

The zip holds every member's personal data and ID documents: copy it to
your own computer, keep it there only, and delete both copies when done.
This is a generated export: SQLite on the server stays the only source of
truth, and editing the spreadsheet changes nothing.
"""
import argparse
import asyncio
import logging
import os
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from telegram import Bot  # noqa: E402
from telegram.error import TelegramError  # noqa: E402

from src import db  # noqa: E402
from src.config import BOT_TOKEN  # noqa: E402
from src.kyc_form import active_fields, display_value, label  # noqa: E402

# Download URLs contain the bot token: never let the HTTP client log them.
logging.getLogger("httpx").setLevel(logging.WARNING)
# ===========================================================================

EXPORT_DIR = ROOT / "data" / "exports"
BASE_COLUMNS = ["No.", "Username", "Telegram name", "Status", "Role",
                "Onboarding lead", "Legacy Member", "First seen", "Last updated"]


def _safe(text):
  return "".join(c if c.isalnum() or c in "-_" else "_" for c in text)[:40] or "member"


def _members(status):
  with db.get_conn() as conn:
    sql = "SELECT * FROM members"
    args = ()
    if status:
      sql += " WHERE status = ?"
      args = (status,)
    sql += " ORDER BY lower(COALESCE(username, first_name, ''))"
    return conn.execute(sql, args).fetchall()


async def _download(bot, file_ref, dest_stem):
  """Download one photo; returns the saved file's name, or None."""
  try:
    f = await bot.get_file(file_ref)
    ext = Path(f.file_path or "").suffix or ".jpg"
    path = dest_stem.with_suffix(ext)
    await f.download_to_drive(path)
    return path.name
  except TelegramError as e:
    print(f"  could not download a photo: {e.__class__.__name__}")
    return None


async def build(status, photos):
  fields = active_fields()
  docs = [f for f in fields if f["type"] == "document"]
  stamp = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H-%M-%SZ")
  work = EXPORT_DIR / f"atl_members_{stamp}"
  (work / "ids").mkdir(parents=True)
  os.chmod(EXPORT_DIR, 0o700)

  wb = Workbook()
  ws = wb.active
  ws.title = "Members"
  headers = BASE_COLUMNS + [label(f) for f in fields] + ["Needs checking", "Awaiting approval", "Expelled"]
  ws.append(headers)
  for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F2937")

  members = _members(status)
  missing_photos = 0
  bot = Bot(BOT_TOKEN) if photos and docs else None
  if bot:
    await bot.initialize()
  try:
    for n, m in enumerate(members, 1):
      uid = m["user_id"]
      s = db.member_summary(uid)
      answers = db.get_kyc_answers(uid)
      pending = sorted({label_for(c["field_key"], fields) for c in db.get_open_changes(uid)})
      who = f"@{m['username']}" if m["username"] else ""
      name = f"{m['first_name'] or ''} {m['last_name'] or ''}".strip()
      row = [n, who, name, m["status"], s["role"] or "member",
             "yes" if s["manager"] else "", "yes" if s["legacy_member"] else "",
             m["first_seen"], m["updated_at"]]
      links = {}
      review = []
      for f in fields:
        a = answers.get(f["key"])
        if a is not None and a["needs_review"]:
          review.append(label(f))
        if f["type"] == "document":
          if a and a["file_ref"]:
            saved = None
            if bot:
              saved = await _download(bot, a["file_ref"],
                                      work / "ids" / f"{n:04d}_{_safe(who or name)}_{f['key']}")
              await asyncio.sleep(0.05)   # stay well inside Telegram's rate limits
            if saved:
              links[len(row)] = f"ids/{saved}"
              row.append("Open photo")
            else:
              row.append("on file, download failed" if photos else "on file")
              if photos:
                missing_photos += 1
          else:
            row.append("")
        else:
          row.append(display_value(f, a["value_text"]) if a else "")
      expelled = s["expelled"]
      row += [", ".join(review), ", ".join(pending),
              f"{expelled['created_at'][:10]}: {expelled['note'] or ''}" if expelled else ""]
      ws.append(row)
      for col, target in links.items():
        cell = ws.cell(row=ws.max_row, column=col + 1)
        cell.hyperlink = target
        cell.style = "Hyperlink"
      if n % 100 == 0:
        print(f"  {n} of {len(members)} members done")
  finally:
    if bot:
      await bot.shutdown()

  ws.freeze_panes = "C2"
  ws.auto_filter.ref = ws.dimensions
  for i, h in enumerate(headers, 1):
    ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(h) + 4))
  wb.save(work / "members.xlsx")

  zip_path = shutil.make_archive(str(work), "zip", root_dir=work)
  shutil.rmtree(work)
  os.chmod(zip_path, 0o600)
  return Path(zip_path), len(members), missing_photos


def label_for(key, fields):
  f = next((f for f in fields if f["key"] == key), None)
  return label(f) if f else key


def main():
  ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
  ap.add_argument("--status", choices=db.ALL_STATUSES, help="only members with this status")
  ap.add_argument("--no-photos", action="store_true", help="skip downloading ID photos")
  args = ap.parse_args()
  print("Exporting…")
  zip_path, count, missing = asyncio.run(build(args.status, not args.no_photos))
  print(f"\nDone: {count} members -> {zip_path.relative_to(ROOT)}")
  if missing:
    print(f"{missing} photos couldn't be downloaded; their cells say so.")
  print("Copy it to your computer, then delete it here:  rm", zip_path.relative_to(ROOT))


if __name__ == "__main__":
  main()
